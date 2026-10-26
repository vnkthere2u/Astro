"""
build_excel.py — generates AstroMarket_Workbook.xlsx from the engine.
Engine outputs (scores, ephemeris) are written as values; every summary, call, colour and
tracker statistic is a live Excel formula / conditional format driven by the Settings sheet.
Re-run `python build_excel.py [start_date] [days]` to roll the workbook forward.
"""
from __future__ import annotations
import json, sys
from datetime import date, datetime, timedelta
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
import astro_core as ac, engine as E, calib as C, tracker as TR

HERE = Path(__file__).parent
F = 'Arial'
H1 = Font(name=F, size=14, bold=True, color='1A365D'); H2 = Font(name=F, size=11, bold=True, color='FFFFFF')
B = Font(name=F, size=10); BB = Font(name=F, size=10, bold=True); BLUE = Font(name=F, size=10, color='0000FF')
GREEN_T = Font(name=F, size=10, color='008000'); SMALL = Font(name=F, size=8, color='666666')
HDR = PatternFill('solid', fgColor='1A365D'); SUB = PatternFill('solid', fgColor='E2E8F0'); YEL = PatternFill('solid', fgColor='FFFF00')
G_F = PatternFill('solid', fgColor='1F9D55'); A_F = PatternFill('solid', fgColor='D69E2E'); R_F = PatternFill('solid', fgColor='C53030')
G_L = PatternFill('solid', fgColor='C6F6D5'); A_L = PatternFill('solid', fgColor='FEFCBF'); R_L = PatternFill('solid', fgColor='FED7D7')
THIN = Border(*(Side(style='thin', color='CBD5E0'),) * 4)
CENTER = Alignment(horizontal='center', vertical='center'); WRAP = Alignment(wrap_text=True, vertical='top')
THR = 'Settings!$C$4'

def header(ws, row, cols, start_col=1):
    for i, c in enumerate(cols):
        x = ws.cell(row=row, column=start_col + i, value=c); x.font = H2; x.fill = HDR; x.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); x.border = THIN

def title(ws, text, sub=None):
    ws['A1'] = text; ws['A1'].font = H1
    if sub: ws['A2'] = sub; ws['A2'].font = SMALL

def rag_cf(ws, rng, light=False):
    gf, af, rf = (G_L, A_L, R_L) if light else (G_F, A_F, R_F)
    first = rng.split(':')[0]
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER({first}),{first}>{THR})'], fill=gf))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER({first}),{first}<-{THR})'], fill=rf))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER({first}),ABS({first})<={THR})'], fill=af))

def widths(ws, w: dict):
    for k, v in w.items(): ws.column_dimensions[k].width = v

def build(start: date, ndays=365, out=None):
    eng = E.Engine(muhurat=[(date(2023, 11, 12), 354), (date(2024, 11, 1), 335), (date(2025, 10, 21), 63)])
    days = [start + timedelta(days=i) for i in range(ndays)]
    res = {d: eng.view(d) for d in days}
    TG = ['NIFTY'] + E.SECTORS + E.COMMODITIES
    wb = Workbook()

    # ------------------------------------------------------------ Settings
    st = wb.active; st.title = 'Settings'
    title(st, 'Settings & legend', 'Yellow cells are inputs. Change them and every call, colour and tracker statistic in the workbook updates.')
    rows = [('Green / Red threshold (|score| above which a call is Bullish / Bearish)', E.THRESH, 'Engine default 0.15. Scores run −1 … +1.'),
            ('Overall-view weight — today', 0.5, 'Overall = w·today + w·next week + w·next 2 months (pre-computed in Daily_Composite with these defaults).'),
            ('Overall-view weight — next week', 0.3, ''), ('Overall-view weight — next 2 months', 0.2, ''),
            ('Workbook generated', datetime.now().strftime('%Y-%m-%d %H:%M'), 'Re-run build_excel.py to roll forward.'),
            ('Market location / opening bell', f"{eng.exchange} {eng.open_t:%H:%M}", 'Opening-bell chart & horā computed for this place.')]
    header(st, 3, ['', 'Parameter', 'Value', 'Note'])
    for i, (a, b, c) in enumerate(rows):
        st.cell(row=4 + i, column=2, value=a).font = B
        x = st.cell(row=4 + i, column=3, value=b); x.font = BLUE if i < 4 else B
        if i < 4: x.fill = YEL
        st.cell(row=4 + i, column=4, value=c).font = SMALL
    st['B12'] = 'Legend'; st['B12'].font = BB
    for i, (f, t) in enumerate([(G_F, 'Bullish (score > threshold)'), (A_F, 'Neutral (|score| ≤ threshold)'), (R_F, 'Bearish (score < −threshold)'),
                                (YEL, 'Input cell (blue text) — edit me')]):
        st.cell(row=13 + i, column=2).fill = f; st.cell(row=13 + i, column=3, value=t).font = B
    st['B18'] = 'Important'; st['B18'].font = BB
    st['B19'] = ('Backtest 2012-02-21 → 2026-10-23 on NSE indices: intraday directional hit-rate 50.5% (95% CI 50.0–51.0%), short-term 50.5%, '
                 'long-term 46.1% vs ~60% for a naive "always up" call. Treat outputs as an implementation of the book, not as a validated trading signal.')
    st['B19'].alignment = WRAP; st['B19'].font = B; st.merge_cells('B19:D22')
    widths(st, {'A': 2, 'B': 62, 'C': 26, 'D': 80})

    # ------------------------------------------------------------ Dashboard
    db = wb.create_sheet('Dashboard', 0)
    d0 = E.is_trading_day(start) and start or next(d for d in days if d.weekday() < 5)
    r0 = res[d0]
    title(db, f'AstroMarket Dashboard — first trading day {d0:%A %d %b %Y}', 'Today = that session · Short = next 5 sessions (averaged) · Long = next 40 sessions (averaged)')
    db['A4'] = 'Sky'; db['A4'].font = BB
    cm = r0.chart_mid
    db['B4'] = f"Candra {ac.fmt_deg(cm['Mo'].lon)} · {ac.NAK[cm.moon_nak]} · opening lagna {ac.fmt_deg(r0.chart_open.asc)}"
    db['A5'] = 'Business cycle'; db['A5'].font = BB; db['B5'] = f"{r0.regime['phase']} [{r0.regime['state']}] — {r0.regime['detail']}"
    db['A6'] = 'Book guidance'; db['A6'].font = BB; db['B6'] = r0.regime['desc']
    header(db, 8, ['Target', 'Intraday', 'Short', 'Long', 'Composite', 'Call', 'Intraday conf.', 'Short conf.', 'Long conf.'])
    for i, t in enumerate(TG):
        rr = 9 + i
        db.cell(row=rr, column=1, value=E.TARGET_NAME[t]).font = B
        for j, hz in enumerate(E.HORIZONS):
            db.cell(row=rr, column=2 + j, value=round(r0.scores[t][hz], 3)).font = B
            cf = C.confidence(t, hz, r0.scores[t][hz])
            x = db.cell(row=rr, column=7 + j, value=cf['p'] if cf else None); x.number_format = '0%'; x.font = B
        db.cell(row=rr, column=5, value=f'=Settings!$C$5*B{rr}+Settings!$C$6*C{rr}+Settings!$C$7*D{rr}').font = B
        db.cell(row=rr, column=6, value=f'=IF(E{rr}>{THR},"Bullish",IF(E{rr}<-{THR},"Bearish","Neutral"))').font = BB
        for c in range(2, 6): db.cell(row=rr, column=c).number_format = '+0.00;-0.00;0.00'
    last = 8 + len(TG)
    rag_cf(db, f'B9:E{last}', light=True)
    db.conditional_formatting.add(f'F9:F{last}', CellIsRule(operator='equal', formula=['"Bullish"'], fill=G_L))
    db.conditional_formatting.add(f'F9:F{last}', CellIsRule(operator='equal', formula=['"Bearish"'], fill=R_L))
    plan = E.action_plan(r0, C.confidence)
    rr = last + 2; db.cell(row=rr, column=1, value='Intraday timing (horā, Ch.10)').font = BB
    for k, w in enumerate(plan['hora_windows']): db.cell(row=rr + 1 + k, column=1, value=w).font = B
    rr += len(plan['hora_windows']) + 2; db.cell(row=rr, column=1, value='Micro-daśā windows (Ch.12)').font = BB
    for k, w in enumerate(plan['dasha_windows']): db.cell(row=rr + 1 + k, column=1, value=w).font = B
    rr += len(plan['dasha_windows']) + 2; db.cell(row=rr, column=1, value='Notes').font = BB
    for k, w in enumerate(plan['notes']): db.cell(row=rr + 1 + k, column=1, value=w).font = B
    db.cell(row=7, column=7, value='Confidence = shrunk historical hit-rate of comparable calls (same target, horizon, strength). ~50% means no demonstrated edge.').font = SMALL
    widths(db, {'A': 34, 'B': 11, 'C': 11, 'D': 11, 'E': 11, 'F': 11, 'G': 13, 'H': 13, 'I': 13})
    db.freeze_panes = 'B9'

    # ------------------------------------------------------------ Daily_Composite (long table, rows = dates)
    dc = wb.create_sheet('Daily_Composite')
    title(dc, 'Daily composite score by target (engine output) — rows = calendar days', 'Composite = 0.5·intraday + 0.3·short + 0.2·long. Colours follow Settings threshold.')
    cols = ['Date', 'Trading day', 'Candra rāśi', 'Nakṣatra', 'Lagna @open', 'Regime'] + [E.TARGET_NAME[t] for t in TG]
    header(dc, 4, cols)
    for i, d in enumerate(days):
        r = res[d]; rr = 5 + i
        dc.cell(row=rr, column=1, value=d).number_format = 'dd-mmm-yy'
        dc.cell(row=rr, column=2, value=1 if d.weekday() < 5 else 0)
        dc.cell(row=rr, column=3, value=ac.RASI[r.chart_mid['Mo'].sign]); dc.cell(row=rr, column=4, value=ac.NAK[r.chart_mid.moon_nak])
        dc.cell(row=rr, column=5, value=ac.RASI[int(r.chart_open.asc // 30)]); dc.cell(row=rr, column=6, value=f"{r.regime['phase']} [{r.regime['state'].split(' ')[0]}]")
        for j, t in enumerate(TG):
            x = dc.cell(row=rr, column=7 + j, value=round(r.composite(t), 3)); x.number_format = '+0.00;-0.00;0.00'
    lastr = 4 + len(days)
    rag_cf(dc, f'G5:{L(6 + len(TG))}{lastr}', light=True)
    for c in range(1, 7 + len(TG)):
        for rr in range(5, lastr + 1): dc.cell(row=rr, column=c).font = B
    widths(dc, {'A': 11, 'B': 8, 'C': 11, 'D': 15, 'E': 11, 'F': 30})
    for j in range(len(TG)): dc.column_dimensions[L(7 + j)].width = 9
    dc.freeze_panes = 'G5'

    # ------------------------------------------------------------ Horizon tables
    for hz in E.HORIZONS:
        ws = wb.create_sheet(f'{hz.capitalize()}_12M')
        title(ws, f'{E.HZ_LABEL[hz]} score by target — rows = calendar days', 'Engine output. Colours follow Settings threshold.')
        header(ws, 4, ['Date'] + [E.TARGET_NAME[t] for t in TG])
        for i, d in enumerate(days):
            ws.cell(row=5 + i, column=1, value=d).number_format = 'dd-mmm-yy'
            for j, t in enumerate(TG):
                x = ws.cell(row=5 + i, column=2 + j, value=round(res[d].scores[t][hz], 3)); x.number_format = '+0.00;-0.00;0.00'; x.font = B
        rag_cf(ws, f'B5:{L(1 + len(TG))}{4 + len(days)}', light=True)
        ws.column_dimensions['A'].width = 11
        for j in range(len(TG)): ws.column_dimensions[L(2 + j)].width = 9
        ws.freeze_panes = 'B5'

    # ------------------------------------------------------------ Grid_12M (day boxes)
    gr = wb.create_sheet('Grid_12M', 2)
    title(gr, 'Day-box grid — 12 months × all sectors & commodities (daily composite)', 'Each box links to Daily_Composite; green = Bullish, amber = Neutral, red = Bearish. Weekend columns shaded headers.')
    gr.cell(row=4, column=1, value='Target').font = BB
    for j, d in enumerate(days):
        col = 2 + j
        if d.day == 1 or j == 0:
            gr.cell(row=3, column=col, value=d.strftime('%b-%y')).font = BB
        x = gr.cell(row=4, column=col, value=d.day); x.font = Font(name=F, size=7, bold=True, color='999999' if d.weekday() >= 5 else '000000'); x.alignment = CENTER
        gr.column_dimensions[L(col)].width = 3.2
    for i, t in enumerate(TG):
        rr = 5 + i
        gr.cell(row=rr, column=1, value=E.TARGET_NAME[t]).font = B
        src_col = L(7 + i)
        for j, d in enumerate(days):
            x = gr.cell(row=rr, column=2 + j, value=f'=Daily_Composite!{src_col}{5 + j}')
            x.font = Font(name=F, size=6, color='FFFFFF'); x.number_format = ';;;'
    rag_cf(gr, f'B5:{L(1 + len(days))}{4 + len(TG)}')
    gr.column_dimensions['A'].width = 30; gr.freeze_panes = 'B5'

    # ------------------------------------------------------------ Monthly_Summary (formulas)
    ms = wb.create_sheet('Monthly_Summary', 3)
    title(ms, 'Monthly summary (formulas over Daily_Composite, trading days only)', 'Avg score, % of trading days Bullish, % Bearish, and the resulting monthly call.')
    months = []
    m = start.replace(day=1)
    while m <= days[-1]:
        months.append(m); m = (m + timedelta(days=32)).replace(day=1)
    hdr = ['Target']
    for m in months: hdr += [f'{m:%b-%y} avg', '% Bullish', '% Bearish']
    header(ms, 4, hdr)
    ms.cell(row=3, column=1, value='Month start').font = SMALL
    for k, m in enumerate(months):
        c0 = 2 + 3 * k
        a = ms.cell(row=3, column=c0, value=max(m, start)); a.number_format = 'dd-mmm-yy'; a.font = SMALL
        e = min((m + timedelta(days=32)).replace(day=1) - timedelta(days=1), days[-1])
        b_ = ms.cell(row=2, column=c0, value=e); b_.number_format = 'dd-mmm-yy'; b_.font = SMALL
    ms.cell(row=2, column=1, value='Month end').font = SMALL
    for i, t in enumerate(TG):
        rr = 5 + i; ms.cell(row=rr, column=1, value=E.TARGET_NAME[t]).font = B
        col = f'Daily_Composite!${L(7 + i)}$5:${L(7 + i)}${lastr}'
        dt = f'Daily_Composite!$A$5:$A${lastr}'; td = f'Daily_Composite!$B$5:$B${lastr}'
        for k, m in enumerate(months):
            c0 = 2 + 3 * k; s_ = f'{L(c0)}$3'; e_ = f'{L(c0)}$2'
            base = f'{dt},">="&{s_},{dt},"<="&{e_},{td},1'
            x = ms.cell(row=rr, column=c0, value=f'=IFERROR(AVERAGEIFS({col},{base}),"")'); x.number_format = '+0.00;-0.00;0.00'
            y = ms.cell(row=rr, column=c0 + 1, value=f'=IFERROR(COUNTIFS({base},{col},">"&{THR})/COUNTIFS({base}),"")'); y.number_format = '0%'
            z = ms.cell(row=rr, column=c0 + 2, value=f'=IFERROR(COUNTIFS({base},{col},"<"&-{THR})/COUNTIFS({base}),"")'); z.number_format = '0%'
            for c in (x, y, z): c.font = B
        for k in range(len(months)):
            c = L(2 + 3 * k); rag_cf(ms, f'{c}5:{c}{4 + len(TG)}', light=True)
    ms.column_dimensions['A'].width = 30
    for c in range(2, 2 + 3 * len(months)): ms.column_dimensions[L(c)].width = 9
    ms.freeze_panes = 'B5'

    # ------------------------------------------------------------ Intraday_Timing (next 22 trading days)
    it = wb.create_sheet('Intraday_Timing')
    title(it, 'Horā (Ch.10) & micro-daśā (Ch.12) windows — next 22 trading days', 'Signal: BUY before/during = lord favourably aspected; SELL during = afflicted.')
    header(it, 4, ['Date', 'Type', 'From', 'To', 'Lord', 'Net aspect', 'Signal', 'Sectors / commodities'])
    rr = 5; tdays = [d for d in days if d.weekday() < 5][:22]
    for d in tdays:
        r = res[d]
        for h in r.hora:
            for c, v in enumerate([d, 'Horā', h['start'].strftime('%H:%M'), h['end'].strftime('%H:%M'), ac.NAMES[h['lord']], h['q'], h['signal'], h['text']]):
                it.cell(row=rr, column=c + 1, value=v).font = B
            it.cell(row=rr, column=1).number_format = 'dd-mmm-yy'; rr += 1
        for x in r.dasha:
            for c, v in enumerate([d, 'Daśā', x['start'].strftime('%H:%M'), x['end'].strftime('%H:%M'), ac.NAMES[x['lord']], x['q'], x['signal'], ', '.join(E.TARGET_NAME.get(t, t) for t in x['targets'])]):
                it.cell(row=rr, column=c + 1, value=v).font = B
            it.cell(row=rr, column=1).number_format = 'dd-mmm-yy'; rr += 1
    rag_cf(it, f'F5:F{rr}', light=True)
    widths(it, {'A': 11, 'B': 7, 'C': 7, 'D': 7, 'E': 10, 'F': 10, 'G': 18, 'H': 110}); it.freeze_panes = 'A5'

    # ------------------------------------------------------------ Ephemeris & transit calendar
    ep = wb.create_sheet('Ephemeris_12M')
    title(ep, 'Sidereal ephemeris (Lahiri, true node) at mid-session IST', 'Source: Swiss Ephemeris via pyswisseph. R = retrograde.')
    cols = ['Date'] + [f'{ac.NAMES[g]}' for g in ac.GRAHAS] + ['Candra nakṣatra', 'Tithi']
    header(ep, 4, cols)
    for i, d in enumerate(days):
        c = res[d].chart_mid; rr = 5 + i
        ep.cell(row=rr, column=1, value=d).number_format = 'dd-mmm-yy'
        for j, g in enumerate(ac.GRAHAS):
            ep.cell(row=rr, column=2 + j, value=ac.fmt_deg(c[g].lon) + (' R' if c[g].retro else '')).font = B
        ep.cell(row=rr, column=2 + len(ac.GRAHAS), value=ac.NAK[c.moon_nak]).font = B
        ep.cell(row=rr, column=3 + len(ac.GRAHAS), value=c.tithi).font = B
    ep.column_dimensions['A'].width = 11
    for j in range(len(ac.GRAHAS) + 2): ep.column_dimensions[L(2 + j)].width = 17
    ep.freeze_panes = 'B5'
    tc = wb.create_sheet('Transit_Calendar')
    title(tc, 'Transit calendar — ingresses, stations, eclipses (next 12 months)')
    header(tc, 4, ['Date', 'Event', 'Book chapter relevance'])
    ev = []
    for g in ['Su', 'Ma', 'Me', 'Ju', 'Ve', 'Sa', 'Ra', 'Ur', 'Ne', 'Pl']:
        for t, a, b in ac.find_ingresses(g, start, days[-1]):
            ev.append((t.date(), f'{ac.NAMES[g]} enters {ac.RASI[b]} (from {ac.RASI[a]})', 'Ch.7 transit table; Ch.6 trigger rāśi' if g in ('Su', 'Ju') else 'Ch.7 transit table'))
    prev = None
    for d in days:
        c = res[d].chart_mid; rs = {g: c[g].retro for g in ['Me', 'Ve', 'Ma', 'Ju', 'Sa', 'Ur', 'Ne', 'Pl']}
        if prev:
            for g in rs:
                if rs[g] != prev[g]: ev.append((d, f"{ac.NAMES[g]} turns {'retrograde' if rs[g] else 'direct'}", 'Ch.7 rules 5-6 (retrograde intensifies)'))
        prev = rs
    for y in {start.year, days[-1].year}:
        for k, t in ac.eclipses_in_year(y):
            dd = ac.from_jd(t).date()
            if start <= dd <= days[-1]: ev.append((dd, f'{k.title()} eclipse', 'Ch.7 rules 9-10; Ch.15 eclipse cycle'))
    pr = None
    for d in days:
        rg = res[d].regime; key = (rg['phase'], rg['state'])
        if key != pr: ev.append((d, f"Business-cycle regime → {rg['phase']} [{rg['state']}]", 'Ch.16')); pr = key
    for i, (d, e, ch) in enumerate(sorted(ev)):
        tc.cell(row=5 + i, column=1, value=d).number_format = 'dd-mmm-yy'; tc.cell(row=5 + i, column=2, value=e).font = B; tc.cell(row=5 + i, column=3, value=ch).font = SMALL
    widths(tc, {'A': 12, 'B': 60, 'C': 45})

    # ------------------------------------------------------------ Tracker (formulas)
    tr = wb.create_sheet('Tracker', 4)
    title(tr, 'Prediction tracker — predicted vs actual', 'Enter the realised % move (as a fraction, e.g. 0.012 = +1.2%) in the yellow column. Calls, hits and accuracy are formulas.')
    tr['A3'] = 'Example: Nifty intraday predicted −0.42 (Bearish); actual open→close −0.6% → enter −0.006 → Hit = 1.'; tr['A3'].font = SMALL
    header(tr, 5, ['Date', 'Target', 'Horizon', 'Predicted score', 'Predicted call', 'Confidence', 'Actual return', 'Actual direction', 'Hit (1/0)', 'Mode'])
    T = TR.load()
    keep = ['NIFTY', 'BANK', 'IT', 'AUTO', 'PHARMA', 'FMCG', 'METAL', 'REALTY', 'MEDIA', 'INFRA', 'PSU', 'GOLD', 'SILVER', 'CRUDE', 'NATGAS']
    T = T[T.target.isin(keep) & T.horizon.isin(['intraday', 'short'])]
    fut = []
    for d in [x for x in days if x.weekday() < 5][:22]:
        for t in keep:
            for hz in ('intraday', 'short'):
                s = res[d].scores[t][hz]; cf = C.confidence(t, hz, s)
                fut.append(dict(date=pd.Timestamp(d), target=t, horizon=hz, score=round(s, 4), confidence=cf['p'] if cf else None, actual_ret=None, mode='forward'))
    T = pd.concat([T[['date', 'target', 'horizon', 'score', 'confidence', 'actual_ret', 'mode']], pd.DataFrame(fut)], ignore_index=True).sort_values(['date', 'target', 'horizon'])
    rr = 6
    for _, x in T.iterrows():
        tr.cell(row=rr, column=1, value=pd.Timestamp(x['date']).date()).number_format = 'dd-mmm-yy'
        tr.cell(row=rr, column=2, value=E.TARGET_NAME[x['target']]); tr.cell(row=rr, column=3, value=x['horizon'])
        tr.cell(row=rr, column=4, value=float(x['score'])).number_format = '+0.00;-0.00;0.00'
        tr.cell(row=rr, column=5, value=f'=IF(D{rr}>{THR},"Bullish",IF(D{rr}<-{THR},"Bearish","Neutral"))')
        c6 = tr.cell(row=rr, column=6, value=None if pd.isna(x['confidence']) else float(x['confidence'])); c6.number_format = '0%'
        g = tr.cell(row=rr, column=7, value=None if pd.isna(x['actual_ret']) else float(x['actual_ret'])); g.number_format = '+0.00%;-0.00%;0.00%'; g.font = BLUE; g.fill = YEL
        tr.cell(row=rr, column=8, value=f'=IF(G{rr}="","",IF(G{rr}>0,"Bullish",IF(G{rr}<0,"Bearish","Flat")))')
        tr.cell(row=rr, column=9, value=f'=IF(OR(G{rr}="",E{rr}="Neutral"),"",IF(E{rr}=H{rr},1,0))')
        tr.cell(row=rr, column=10, value=x['mode'])
        for c in range(1, 11):
            if c != 7: tr.cell(row=rr, column=c).font = B
        rr += 1
    tl = rr - 1
    rag_cf(tr, f'D6:D{tl}', light=True)
    # summary block
    tr['L5'] = 'Accuracy summary'; tr['L5'].font = BB
    lines = [('All settled calls', f'=IFERROR(AVERAGE(I6:I{tl}),"")', f'=COUNT(I6:I{tl})'),
             ('Intraday', f'=IFERROR(AVERAGEIFS(I6:I{tl},C6:C{tl},"intraday"),"")', f'=COUNTIFS(C6:C{tl},"intraday",I6:I{tl},">=0")'),
             ('Short-term (5 sessions)', f'=IFERROR(AVERAGEIFS(I6:I{tl},C6:C{tl},"short"),"")', f'=COUNTIFS(C6:C{tl},"short",I6:I{tl},">=0")'),
             ('Bullish calls', f'=IFERROR(AVERAGEIFS(I6:I{tl},E6:E{tl},"Bullish"),"")', f'=COUNTIFS(E6:E{tl},"Bullish",I6:I{tl},">=0")'),
             ('Bearish calls', f'=IFERROR(AVERAGEIFS(I6:I{tl},E6:E{tl},"Bearish"),"")', f'=COUNTIFS(E6:E{tl},"Bearish",I6:I{tl},">=0")'),
             ('Benchmark: always Bullish, same rows', f'=IFERROR(COUNTIFS(H6:H{tl},"Bullish",I6:I{tl},">=0")/COUNT(I6:I{tl}),"")', '')]
    header(tr, 6, ['Slice', 'Hit-rate', 'n'], start_col=12)
    for i, (a, b_, c) in enumerate(lines):
        tr.cell(row=7 + i, column=12, value=a).font = B
        x = tr.cell(row=7 + i, column=13, value=b_); x.number_format = '0.0%'; x.font = BB
        tr.cell(row=7 + i, column=14, value=c).font = B
    tr['L14'] = 'Per target (all horizons)'; tr['L14'].font = BB
    header(tr, 15, ['Target', 'Hit-rate', 'n'], start_col=12)
    for i, t in enumerate(keep):
        nm = E.TARGET_NAME[t]
        tr.cell(row=16 + i, column=12, value=nm).font = B
        x = tr.cell(row=16 + i, column=13, value=f'=IFERROR(AVERAGEIFS($I$6:$I${tl},$B$6:$B${tl},L{16 + i}),"")'); x.number_format = '0.0%'
        tr.cell(row=16 + i, column=14, value=f'=COUNTIFS($B$6:$B${tl},L{16 + i},$I$6:$I${tl},">=0")')
    widths(tr, {'A': 11, 'B': 26, 'C': 9, 'D': 10, 'E': 10, 'F': 10, 'G': 11, 'H': 11, 'I': 8, 'J': 9, 'K': 2, 'L': 30, 'M': 10, 'N': 7})
    tr.freeze_panes = 'A6'
    dv = DataValidation(type='decimal', operator='between', formula1='-1', formula2='1', allow_blank=True); tr.add_data_validation(dv); dv.add(f'G6:G{tl}')

    # ------------------------------------------------------------ Rulebook sheets
    RB = E.RB
    ws = wb.create_sheet('RB_Ch6_Sectors'); title(ws, 'Ch.6 — sector rules: 30 sectors × 3 horizons × buy/sell (180 statements)')
    header(ws, 4, ['Sector', 'Horizon', 'Side', 'Kāraka grahas', 'Trigger graha', 'Trigger rāśis', 'Trigger nakṣatras', 'Condition', 'Outcome', 'Page', 'Editor note'])
    rr = 5
    for r in RB['ch6']:
        for hz, sd in r['horizons'].items():
            for side, x in sd.items():
                vals = [r['name'], hz, side, ', '.join(ac.NAMES[k] for k in x['karakas']), ac.NAMES[x['trigger_planet']],
                        ', '.join(ac.RASI[i] for i in x['trigger_rasis']), ', '.join(ac.NAK[i] for i in x['trigger_naks']),
                        f"kārakas aspected by {'benefic' if side == 'buy' else 'malefic'} grahas while trigger graha in listed rāśi/nakṣatra",
                        'BUYING emerges' if side == 'buy' else 'SELLING emerges', x['page'], ' '.join(x['editor_note'])]
                for c, v in enumerate(vals): ws.cell(row=rr, column=c + 1, value=v).font = B
                rr += 1
    widths(ws, {'A': 26, 'B': 9, 'C': 6, 'D': 34, 'E': 10, 'F': 50, 'G': 70, 'H': 50, 'I': 16, 'J': 6, 'K': 20}); ws.freeze_panes = 'A5'
    ws = wb.create_sheet('RB_Ch7_Transits'); title(ws, 'Ch.7 — transit of 12 grahas × 12 rāśis with aspect conditions (all combinations)')
    header(ws, 4, ['Graha', 'Rāśi', 'Condition A: aspected by', 'Outcome A (−1…+1)', 'Condition B: aspected by', 'Outcome B', 'If no qualifying aspect', 'Sectors / commodities affected', 'Horizon', 'Page', 'Fix', 'Book text'])
    for i, e in enumerate(RB['ch7']):
        w = e['whereas'] or {}
        vals = [ac.NAMES[e['graha']], ac.RASI[e['rasi']], ', '.join(ac.NAMES[a] for a in e['cond_aspectors']) or '(unconditional)', e['effect'],
                ', '.join(ac.NAMES[a] for a in w.get('aspectors', [])), w.get('effect'), 'Rule 1: dignity × 0.3' if not e['unconditional'] else 'unconditional',
                ', '.join(E.TARGET_NAME.get(t, t) for t in e['targets']), E.CH7_HZ[e['graha']], e['page'], e.get('editor_fix', ''), e.get('text', '')[:500]]
        for c, v in enumerate(vals): ws.cell(row=5 + i, column=c + 1, value=v).font = B
    rag_cf(ws, f'D5:D{4 + len(RB["ch7"])}', light=True); rag_cf(ws, f'F5:F{4 + len(RB["ch7"])}', light=True)
    widths(ws, {'A': 10, 'B': 10, 'C': 26, 'D': 10, 'E': 26, 'F': 10, 'G': 18, 'H': 70, 'I': 9, 'J': 6, 'K': 30, 'L': 100}); ws.freeze_panes = 'C5'
    ws = wb.create_sheet('RB_Ch8_Nakshatra'); title(ws, 'Ch.8 — intraday rules for Candra in each of 27 nakṣatras')
    header(ws, 4, ['Nakṣatra', 'Lord', 'POSITIVE if lord aspected by', 'Stronger if Candra aspected by', 'NEGATIVE if lord aspected by', 'Stronger if Candra aspected by', 'Horā of peak effect', 'Sectors / commodities'])
    for i, n in enumerate(RB['ch8']):
        vals = [n['nakshatra'], ac.NAMES[n['lord']], ', '.join(ac.NAMES[x] for x in n['pos']['aspectors']), ', '.join(ac.NAMES[x] for x in n['pos']['moon_confirm']),
                ', '.join(ac.NAMES[x] for x in n['neg']['aspectors']), ', '.join(ac.NAMES[x] for x in n['neg']['moon_confirm']), ' & '.join(ac.NAMES[x] for x in n['hora']), n['pos']['items']]
        for c, v in enumerate(vals): ws.cell(row=5 + i, column=c + 1, value=v).font = B
    widths(ws, {'A': 16, 'B': 9, 'C': 24, 'D': 24, 'E': 24, 'F': 24, 'G': 16, 'H': 110})
    ws = wb.create_sheet('RB_Ch10_Hora'); title(ws, 'Ch.10 — horā strategies (Fig.1 sequence from sunrise)')
    header(ws, 4, ['Horā lord', 'BUY before horā if aspected by', 'SELL during horā if aspected by', 'Stocks / commodities', 'Page'])
    for i, (g, v) in enumerate(RB['ch10'].items()):
        for c, x in enumerate([ac.NAMES[g], ', '.join(ac.NAMES[a] for a in v['good']), ', '.join(ac.NAMES[a] for a in v['bad']), v['text'], v['page']]):
            ws.cell(row=5 + i, column=c + 1, value=x).font = B
    header(ws, 14, ['Horā #'] + ['Sun', 'Mon', 'Tue', 'Wed', 'Thu', 'Fri', 'Sat'])
    for k in range(24):
        ws.cell(row=15 + k, column=1, value=k + 1).font = B
        for j, wd in enumerate([6, 0, 1, 2, 3, 4, 5]):
            first = ac.WEEKDAY_LORD[wd]; i0 = ac.HORA_CYCLE.index(first)
            ws.cell(row=15 + k, column=2 + j, value=ac.NAMES[ac.HORA_CYCLE[(i0 + k) % 7]]).font = B
    widths(ws, {'A': 12, 'B': 28, 'C': 28, 'D': 110, 'E': 8})
    ws = wb.create_sheet('RB_Ch14_OpeningBell'); title(ws, 'Ch.13/14 — 12 grahas × 12 bhāvas of the opening-bell chart (144 rules)')
    header(ws, 4, ['Graha', 'Bhāva', 'Bhāva signification', 'If aspected by', 'Outcome', 'If afflicted by', 'Outcome', 'Unconditional', 'Mixed effect', 'Sectors', 'Bhāva weight'])
    for i, e in enumerate(RB['ch14']):
        vals = [ac.NAMES[e['graha']], e['house'], e['house_text'], ', '.join(ac.NAMES[x] for x in e['good']), e['effect'], ', '.join(ac.NAMES[x] for x in e['bad']), e['bad_effect'],
                e['unconditional'], json.dumps(e['mixed']) if e['mixed'] else '', ', '.join(E.TARGET_NAME.get(x, x) for x in e['targets']), e['house_weight']]
        for c, v in enumerate(vals): ws.cell(row=5 + i, column=c + 1, value=v).font = B
    for col in ('E', 'G', 'H'): rag_cf(ws, f'{col}5:{col}{4 + len(RB["ch14"])}', light=True)
    widths(ws, {'A': 10, 'B': 6, 'C': 50, 'D': 24, 'E': 9, 'F': 24, 'G': 9, 'H': 10, 'I': 30, 'J': 70, 'K': 8})
    ws = wb.create_sheet('RB_Ch4_5_Grahas_Rasis'); title(ws, 'Ch.4 grahas & Ch.5 rāśis — significations used to map rules to sectors')
    header(ws, 4, ['Graha', 'Own', 'Exalted', 'Debilitated', 'Friends', 'Enemies', 'Sectors', 'Commodities'])
    for i, (g, d) in enumerate(RB['ch4'].items()):
        for c, v in enumerate([ac.NAMES[g], ', '.join(ac.RASI[x] for x in d['own']), ac.RASI[d['exalt']], ac.RASI[d['debil']], ', '.join(ac.NAMES[x] for x in d['friends']),
                               ', '.join(ac.NAMES[x] for x in d['enemies']), d['sectors'], d['commodities']]):
            ws.cell(row=5 + i, column=c + 1, value=v).font = B
    header(ws, 19, ['Rāśi', 'Ruler', 'Element', 'Mode', 'Benefics', 'Malefics', 'Market sectors', ''])
    for i, v in enumerate(RB['ch5'].values()):
        for c, x in enumerate([v['rasi'], ac.NAMES[v['ruler']], v['element'], v['mode'], ', '.join(ac.NAMES[a] for a in v['benefics']), ', '.join(ac.NAMES[a] for a in v['malefics']), v['text']]):
            ws.cell(row=20 + i, column=c + 1, value=x).font = B
    widths(ws, {'A': 12, 'B': 14, 'C': 12, 'D': 12, 'E': 22, 'F': 22, 'G': 100, 'H': 70})

    # ------------------------------------------------------------ Historical evidence (book) + verification
    hv = wb.create_sheet('Historical_Evidence'); title(hv, "Book's historical correlations (backup) with engine verification")
    header(hv, 4, ['Chapter', 'Book claim', 'Engine check', 'Status', 'Page'])
    V = json.loads((HERE / 'data' / 'validation_report.json').read_text())
    for i, t in enumerate(V['tests']):
        for c, x in enumerate([t['group'], f"{t['test']}: {t['book']}", t['engine'] + (f" — {t['note']}" if t.get('note') else ''), 'PASS' if t['pass_'] else 'FAIL', t['page']]):
            hv.cell(row=5 + i, column=c + 1, value=x).font = B
    n = 5 + len(V['tests'])
    hv.conditional_formatting.add(f'D5:D{n}', CellIsRule(operator='equal', formula=['"PASS"'], fill=G_L))
    hv.conditional_formatting.add(f'D5:D{n}', CellIsRule(operator='equal', formula=['"FAIL"'], fill=R_L))
    extra = [('Ch.19', 'Muhurat 09-Nov-2007 −151 pts → 2008 crash', 'User input series in Settings of the app', 'INFO', 90),
             ('Ch.19', 'Muhurat 28-Oct-2008 +499 pts → 2009 recovery', '', 'INFO', 90),
             ('Ch.18', 'Śani + Rāhu in Tulā Jan-2013–Jul-2014 → crude "to cross $250"', 'Conjunction verified (T27); WTI actually fell from ~$105 (Jun-2014) to ~$45 (Jan-2015)', 'PREDICTION FAILED', 90),
             ('Ch.16', 'Watch-list recession 2017-18 (Śani Dhanus, Guru Tulā)', 'Placement verified; S&P 500 +19% in 2017, −6% in 2018; Nifty +29% 2017, +3% 2018', 'PREDICTION NOT BORNE OUT', 86)]
    for i, row in enumerate(extra):
        for c, x in enumerate(row): hv.cell(row=n + 1 + i, column=c + 1, value=x).font = B
    widths(hv, {'A': 14, 'B': 70, 'C': 80, 'D': 22, 'E': 6})

    # ------------------------------------------------------------ Validation & Backtest
    va = wb.create_sheet('Source_Validation'); title(va, 'Ephemeris configuration test matrix (36 hard tests from the book)')
    header(va, 4, ['Ayanāṃśa', 'Node', 'Hard pass', 'Hard total', 'Soft pass', 'Soft total', 'Selected'])
    for i, m in enumerate(V['matrix']):
        for c, x in enumerate([m['ayanamsa'], m['node'], m['hard_pass'], m['hard_total'], m['soft_pass'], m['soft_total'],
                               'YES' if (m['ayanamsa'], m['node']) == (V['chosen']['ayanamsa'], V['chosen']['node']) else '']):
            va.cell(row=5 + i, column=c + 1, value=x).font = B
    f = V['aspect_fit']
    va['A17'] = f"Aspect rule vs book worked examples: recall {f['recall']:.0%}, precision {f['precision']:.0%}; theoretical minimum errors {f['minimum_possible_errors']} due to:"; va['A17'].font = BB
    for i, c in enumerate(f['contradictions']): va.cell(row=18 + i, column=1, value='• ' + c).font = B
    widths(va, {'A': 16, 'B': 8, 'C': 10, 'D': 10, 'E': 10, 'F': 10, 'G': 10})
    bt = wb.create_sheet('Backtest'); cal = C.cal()
    title(bt, f"Backtest {cal['meta'].get('period')} — directional hit-rates (threshold ±{cal['meta']['threshold']})",
          'Hit = sign(score) equals sign(realised return). Naive = always predicting the majority direction on the same days.')
    header(bt, 4, ['Target', 'Price series', 'Quality', 'Horizon', 'Calls (n)', 'Hit-rate', '95% low', '95% high', 'Naive', 'In-sample 2012-19', 'Out-of-sample 2020+', 'Avg ret Bullish calls', 'Avg ret Bearish calls'])
    rr = 5
    for t, info in cal['targets'].items():
        for hz, H in info['horizons'].items():
            a = H['all']
            vals = [E.TARGET_NAME.get(t, t), info['meta']['key'], info['meta']['quality'], hz, a['n'], a['hit'], a['lo'], a['hi'], a['naive_hit'],
                    H['in_sample']['hit'], H['out_of_sample']['hit'], a['green_mean'], a['red_mean']]
            for c, x in enumerate(vals):
                if isinstance(x, str) and x in ('nan', 'None'): x = None
                if isinstance(x, float) and x != x: x = None
                cell = bt.cell(row=rr, column=c + 1, value=x); cell.font = B
                if c in (5, 6, 7, 8, 9, 10): cell.number_format = '0.0%'
                if c in (11, 12): cell.number_format = '+0.00%;-0.00%'
            rr += 1
    bt.conditional_formatting.add(f'F5:F{rr}', FormulaRule(formula=['AND(ISNUMBER(F5),G5>0.5)'], fill=G_L))
    bt.conditional_formatting.add(f'F5:F{rr}', FormulaRule(formula=['AND(ISNUMBER(F5),H5<0.5)'], fill=R_L))
    rr += 1; bt.cell(row=rr, column=1, value='By rule family (direct-quality series)').font = BB; rr += 1
    header(bt, rr, ['Family|horizon', 'n', 'Hit-rate', '95% low', '95% high', 'p-value']); rr += 1
    for k, v in sorted(cal['family'].items()):
        for c, x in enumerate([k, v['n'], v['hit'], v['lo'], v['hi'], v['p_value']]):
            cell = bt.cell(row=rr, column=c + 1, value=x); cell.font = B
            if c in (2, 3, 4): cell.number_format = '0.0%'
            if c == 5: cell.number_format = '0.000'
        rr += 1
    rr += 1; bt.cell(row=rr, column=1, value='Long-run cycle tests').font = BB; rr += 1
    for k, v in cal.get('longrun', {}).items():
        bt.cell(row=rr, column=1, value=k).font = BB
        bt.cell(row=rr, column=2, value=json.dumps({x: (round(y, 4) if isinstance(y, float) else y) for x, y in v.items() if x != 'table'})[:1000]).font = SMALL
        rr += 1
        if 'table' in v:
            for kk, vv in v['table'].items():
                bt.cell(row=rr, column=2, value=kk).font = B; bt.cell(row=rr, column=3, value=vv['months']).font = B
                c1 = bt.cell(row=rr, column=4, value=vv['mean_f12']); c1.number_format = '+0.0%'
                c2 = bt.cell(row=rr, column=5, value=vv['pct_up']); c2.number_format = '0%'; rr += 1
    widths(bt, {'A': 26, 'B': 30, 'C': 9, 'D': 9, 'E': 9, 'F': 9, 'G': 9, 'H': 9, 'I': 9, 'J': 12, 'K': 12, 'L': 12, 'M': 12}); bt.freeze_panes = 'A5'

    # order & save
    order = ['Dashboard', 'Settings', 'Grid_12M', 'Monthly_Summary', 'Tracker', 'Daily_Composite', 'Intraday_12M', 'Short_12M', 'Long_12M', 'Intraday_Timing',
             'Transit_Calendar', 'Ephemeris_12M', 'RB_Ch6_Sectors', 'RB_Ch7_Transits', 'RB_Ch8_Nakshatra', 'RB_Ch10_Hora', 'RB_Ch14_OpeningBell', 'RB_Ch4_5_Grahas_Rasis',
             'Historical_Evidence', 'Source_Validation', 'Backtest']
    wb._sheets = [wb[n] for n in order]
    for ws in wb.worksheets: ws.sheet_view.showGridLines = False if ws.title in ('Dashboard', 'Settings') else True
    out = out or (HERE / 'AstroMarket_Workbook.xlsx')
    wb.save(out); return out

if __name__ == '__main__':
    s = date.fromisoformat(sys.argv[1]) if len(sys.argv) > 1 else date.today() + timedelta(days=1)
    n = int(sys.argv[2]) if len(sys.argv) > 2 else 365
    print(build(s, n))
